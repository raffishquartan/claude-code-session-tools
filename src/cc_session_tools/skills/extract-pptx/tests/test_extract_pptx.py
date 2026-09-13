"""Tests for ../scripts/extract_pptx.py.

Exercised end-to-end via subprocess, matching the archive-correspondence skill's test
convention, so the real CLI behaviour (argv parsing, exit codes, stderr messages) is what
we assert on. Fixtures are real .pptx files built at test time with python-pptx (the same
library the script itself depends on) rather than checked-in binary files.
"""
from __future__ import annotations

import io
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

SCRIPT = Path(__file__).resolve().parent.parent / "scripts" / "extract_pptx.py"

HAS_SOFFICE = bool(shutil.which("soffice") or shutil.which("libreoffice"))
HAS_TESSERACT = bool(shutil.which("tesseract"))


def _blank_png_bytes() -> bytes:
    """A real, valid 32x32 white PNG - enough to exercise the OCR code path
    without asserting on recognised text (accuracy on a blank image is not
    something to pin down). Built with Pillow (already an --ocr dependency
    of the script under test) rather than a hand-rolled byte literal."""
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (32, 32), "white").save(buf, format="PNG")
    return buf.getvalue()


def _run(*args: str, timeout: float = 30) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(SCRIPT), *args],
        capture_output=True,
        text=True,
        timeout=timeout,
    )


def _build_pptx(path: Path, *, with_embed: bool = False) -> None:
    """A two-slide deck: slide 1 has a title, a bullet body, a table, and
    speaker notes; slide 2 is intentionally blank (no extractable text)."""
    prs = Presentation()

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly Review"
    body = slide.placeholders[1]
    body.text_frame.text = "First bullet"
    body.text_frame.add_paragraph().text = "Second bullet"

    table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1))
    table = table_shape.table
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "EMEA"
    table.cell(1, 1).text = "42"

    slide.notes_slide.notes_text_frame.text = "Remember to mention the forecast."

    if with_embed:
        slide.shapes.add_ole_object(
            io.BytesIO(b"embedded workbook bytes"),
            prog_id="Package",
            left=Inches(5),
            top=Inches(3),
            width=Inches(1),
            height=Inches(1),
        )

    prs.slides.add_slide(prs.slide_layouts[6])  # blank layout, no shapes

    prs.save(str(path))


def _build_pptx_with_image(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    img_path = Path(tempfile.mkdtemp()) / "blank.png"
    img_path.write_bytes(_blank_png_bytes())
    slide.shapes.add_picture(str(img_path), Inches(1), Inches(1), Inches(1), Inches(1))
    prs.save(str(path))


class ExtractPptxCli(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="extract_pptx_test_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_missing_source_file_exits_1(self) -> None:
        result = _run(str(self.tmp / "nope.pptx"), "--out-md", str(self.tmp / "out.md"))
        self.assertEqual(result.returncode, 1)
        self.assertIn("does not exist", result.stderr)

    def test_happy_path_text_table_notes_no_ocr(self) -> None:
        src = self.tmp / "deck.pptx"
        _build_pptx(src)
        out_md = self.tmp / "out.md"
        result = _run(str(src), "--out-md", str(out_md), "--no-ocr")
        self.assertEqual(result.returncode, 0, result.stderr)
        text = out_md.read_text()
        self.assertIn("Quarterly Review", text)
        self.assertIn("First bullet", text)
        self.assertIn("Second bullet", text)
        self.assertIn("Region | Revenue", text)
        self.assertIn("EMEA | 42", text)
        self.assertIn("Remember to mention the forecast.", text)

    def test_slide_with_no_text_shows_placeholder(self) -> None:
        src = self.tmp / "deck.pptx"
        _build_pptx(src)
        out_md = self.tmp / "out.md"
        _run(str(src), "--out-md", str(out_md), "--no-ocr")
        text = out_md.read_text()
        self.assertIn("*(no extractable text on this slide)*", text)

    def test_embeds_extracted_to_dir_and_listed(self) -> None:
        src = self.tmp / "deck.pptx"
        _build_pptx(src, with_embed=True)
        out_md = self.tmp / "out.md"
        embeds_dir = self.tmp / "embeds"
        result = _run(
            str(src), "--out-md", str(out_md), "--no-ocr", "--embeds-dir", str(embeds_dir)
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        produced = list(embeds_dir.glob("slide001-*"))
        self.assertEqual(len(produced), 1)
        self.assertEqual(produced[0].read_bytes(), b"embedded workbook bytes")
        self.assertIn("## Embedded file objects", out_md.read_text())

    def test_embeds_dir_not_created_when_no_embeds(self) -> None:
        src = self.tmp / "deck.pptx"
        _build_pptx(src, with_embed=False)
        out_md = self.tmp / "out.md"
        embeds_dir = self.tmp / "embeds"
        result = _run(
            str(src), "--out-md", str(out_md), "--no-ocr", "--embeds-dir", str(embeds_dir)
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertFalse(embeds_dir.exists())

    @unittest.skipUnless(HAS_TESSERACT, "tesseract binary not on PATH")
    def test_ocr_enabled_runs_without_crashing_on_blank_image(self) -> None:
        src = self.tmp / "deck.pptx"
        _build_pptx_with_image(src)
        out_md = self.tmp / "out.md"
        result = _run(str(src), "--out-md", str(out_md))
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_pdf_out_without_soffice_exits_1(self) -> None:
        if HAS_SOFFICE:
            self.skipTest("soffice/libreoffice is installed on this machine")
        src = self.tmp / "deck.pptx"
        _build_pptx(src)
        result = _run(
            str(src), "--out-md", str(self.tmp / "out.md"),
            "--no-ocr", "--pdf-out", str(self.tmp / "out.pdf"),
        )
        self.assertEqual(result.returncode, 1)
        self.assertIn("soffice", result.stderr)

    def test_creates_missing_out_md_parent_directory(self) -> None:
        src = self.tmp / "deck.pptx"
        _build_pptx(src)
        out_md = self.tmp / "deeper" / "still" / "out.md"
        result = _run(str(src), "--out-md", str(out_md), "--no-ocr")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(out_md.exists())


if __name__ == "__main__":
    unittest.main()
